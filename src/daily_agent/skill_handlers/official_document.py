"""红头公文 skill handler。

职责
----
模型只负责填槽位（标题、红头、正文等），本 handler 用 python-docx 按公文版式
写成 Word，并落到工作区 ``outputs/``。本轮 ask 结束时
``register_run_deliverables_quietly`` 会自动登记 Artifact，对话页即可预览下载。

版式依据常见党政机关公文做法（红头、武文线、3 号仿宋正文、一级标题黑体），
不追求印刷级 GB/T 9704 全要素；目标是下载后用 Word/WPS 打开即像红头文件。
"""

from __future__ import annotations

import re
from datetime import date
from pathlib import Path
from typing import Any

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from daily_agent.artifact_deliverables import OUTPUTS_DIRNAME
from daily_agent.errors import DailyAgentError
from daily_agent.skill_runtime import (
    SKILL_STATUS_ERROR,
    SKILL_STATUS_SUCCESS,
    SkillExecutionContext,
    SkillRunResult,
)

RED = RGBColor(0xFF, 0x00, 0x00)
BLACK = RGBColor(0x00, 0x00, 0x00)
# 中文字体名写入 OOXML，本机没有对应字体时 Word/WPS 会回退，但不影响文件结构。
FONT_XIAOBIAO = "方正小标宋简体"
FONT_FANGSONG = "仿宋_GB2312"
FONT_HEITI = "黑体"
# 「一、」这类短行视为一级标题；过长则仍当正文，避免把整段通知排成黑体。
_LEVEL1_HEADING_RE = re.compile(r"^[一二三四五六七八九十百]+、.{1,40}$")
_ISO_DATE_RE = re.compile(r"^(\d{4})-(\d{1,2})-(\d{1,2})$")
_UNSAFE_FILENAME_RE = re.compile(r'[\\/:*?"<>|\x00-\x1f]')
_MAX_FILENAME_STEM = 80


def run(*, context: SkillExecutionContext) -> SkillRunResult:
    """根据槽位生成红头 Word，写入 outputs/。"""
    inputs = dict(context.inputs)
    try:
        payload = _normalize_inputs(inputs)
        relative_path = _write_docx(payload, working_directory=context.working_directory)
        title = payload["title"]
        response = (
            f"已生成红头公文《{title}》，文件在 {relative_path}。"
            "对话里会出现产物卡片，可以直接预览和下载，不用去代码目录找。"
        )
        return SkillRunResult(
            status=SKILL_STATUS_SUCCESS,
            response_text=response,
            stop_reason="红头公文已生成",
            inputs=payload,
            output_text=response,
            observation={
                "path": relative_path,
                "title": title,
                "filename": Path(relative_path).name,
            },
            clear_session=True,
        )
    except DailyAgentError as exc:
        return SkillRunResult(
            status=SKILL_STATUS_ERROR,
            response_text=str(exc),
            stop_reason="红头公文生成失败",
            inputs=inputs,
            clear_session=True,
        )
    except Exception as exc:
        return SkillRunResult(
            status=SKILL_STATUS_ERROR,
            response_text=f"红头公文生成失败：{exc}",
            stop_reason="红头公文生成失败",
            inputs=inputs,
            observation={"error": type(exc).__name__},
            clear_session=True,
        )


def _normalize_inputs(inputs: dict[str, Any]) -> dict[str, str]:
    """把工具入参收成渲染所需的纯字符串槽位。"""
    title = _as_text(inputs.get("title"))
    body = _as_text(inputs.get("body"))
    red_header = _as_text(inputs.get("red_header"))
    if not title:
        raise DailyAgentError("还缺公文标题，例如「关于××的通知」。")
    if not body:
        raise DailyAgentError("还缺公文正文。请按「一、」「（一）」整理成纯文本后再生成。")
    if not red_header:
        raise DailyAgentError("还缺红头（发文机关版头），例如「××公司文件」。请先问用户是哪个单位发文。")

    recipients = _as_text(inputs.get("recipients"))
    if recipients and not recipients.endswith(("：", ":")):
        recipients = f"{recipients}："

    issuer = _as_text(inputs.get("issuer")) or _issuer_from_red_header(red_header)
    issue_date = _normalize_issue_date(_as_text(inputs.get("issue_date")))
    filename = _safe_filename(_as_text(inputs.get("filename")) or title)
    return {
        "title": title,
        "body": body,
        "red_header": red_header,
        "doc_number": _as_text(inputs.get("doc_number")),
        "recipients": recipients,
        "issuer": issuer,
        "issue_date": issue_date,
        "filename": filename,
    }


def _write_docx(payload: dict[str, str], *, working_directory: str) -> str:
    """渲染并保存 docx，返回相对工作区的路径（POSIX）。"""
    workspace = Path(working_directory).resolve()
    outputs = workspace / OUTPUTS_DIRNAME
    outputs.mkdir(parents=True, exist_ok=True)
    target = (outputs / payload["filename"]).resolve()
    if not _is_relative_to(target, outputs):
        raise DailyAgentError("输出路径不合法，文件名不能包含路径。")

    document = _build_document(payload)
    document.save(str(target))
    return target.relative_to(workspace).as_posix()


def _build_document(payload: dict[str, str]) -> Document:
    """按红头公文版式组装 Document，不落盘。"""
    document = Document()
    section = document.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    # 常用公文页边距：上 3.7、下 3.5、左 2.8、右 2.6 厘米。
    section.top_margin = Cm(3.7)
    section.bottom_margin = Cm(3.5)
    section.left_margin = Cm(2.8)
    section.right_margin = Cm(2.6)

    _add_para(
        document,
        payload["red_header"],
        FONT_XIAOBIAO,
        40,
        align=WD_ALIGN_PARAGRAPH.CENTER,
        line_pt=None,
        color=RED,
        space_after=8,
        run_spacing=600,
    )
    # 空段只用来挂红色底边，充当红头下的武文线。
    _add_para(
        document,
        "",
        FONT_FANGSONG,
        12,
        line_pt=12,
        red_bottom_border=True,
        space_after=0,
    )
    if payload["doc_number"]:
        _add_para(
            document,
            payload["doc_number"],
            FONT_FANGSONG,
            16,
            align=WD_ALIGN_PARAGRAPH.CENTER,
            space_before=8,
            space_after=20,
        )
    _add_para(
        document,
        payload["title"],
        FONT_XIAOBIAO,
        22,
        align=WD_ALIGN_PARAGRAPH.CENTER,
        space_before=8 if not payload["doc_number"] else 0,
        space_after=20,
    )
    if payload["recipients"]:
        _add_para(document, payload["recipients"], FONT_FANGSONG, 16, space_after=6)

    for kind, text in _iter_body_blocks(payload["body"]):
        if kind == "heading":
            _add_para(document, text, FONT_HEITI, 16, indent_chars=2)
        else:
            _add_para(document, text, FONT_FANGSONG, 16, indent_chars=2)

    if payload["issuer"]:
        _add_para(
            document,
            payload["issuer"],
            FONT_FANGSONG,
            16,
            align=WD_ALIGN_PARAGRAPH.RIGHT,
            space_before=28,
        )
    if payload["issue_date"]:
        _add_para(
            document,
            payload["issue_date"],
            FONT_FANGSONG,
            16,
            align=WD_ALIGN_PARAGRAPH.RIGHT,
            right_chars=2,
        )
    return document


def _iter_body_blocks(body: str) -> list[tuple[str, str]]:
    """把正文拆成 (heading|body, 文本)。空行分段；一级标题单独成段。"""
    blocks: list[tuple[str, str]] = []
    for chunk in re.split(r"\n\s*\n", body.replace("\r\n", "\n").strip()):
        for line in chunk.split("\n"):
            text = line.strip()
            if not text:
                continue
            kind = "heading" if _LEVEL1_HEADING_RE.match(text) else "body"
            blocks.append((kind, text))
    if not blocks:
        raise DailyAgentError("正文拆开后没有可用段落。")
    return blocks


def _add_para(
    document: Document,
    text: str,
    font_cn: str,
    size_pt: int,
    *,
    align=None,
    indent_chars: int = 0,
    line_pt: int | None = 28,
    bold: bool = False,
    color=BLACK,
    space_before: int = 0,
    space_after: int = 0,
    red_bottom_border: bool = False,
    run_spacing: int | None = None,
    right_chars: int = 0,
):
    """追加一个段落并写好中文字体、行距、缩进和可选红线。"""
    paragraph = document.add_paragraph()
    fmt = paragraph.paragraph_format
    if line_pt:
        fmt.line_spacing = Pt(line_pt)
    fmt.space_before = Pt(space_before)
    fmt.space_after = Pt(space_after)
    if align is not None:
        fmt.alignment = align
    p_pr = paragraph._p.get_or_add_pPr()
    if indent_chars or right_chars:
        ind = p_pr.find(qn("w:ind"))
        if ind is None:
            ind = OxmlElement("w:ind")
            p_pr.append(ind)
        if indent_chars:
            # firstLineChars 以 100 为一字符；firstLine 用 twip（磅×20）。
            ind.set(qn("w:firstLineChars"), str(indent_chars * 100))
            ind.set(qn("w:firstLine"), str(int(size_pt * 20 * indent_chars)))
        if right_chars:
            ind.set(qn("w:rightChars"), str(right_chars * 100))
            ind.set(qn("w:right"), str(int(size_pt * 20 * right_chars)))
    if red_bottom_border:
        border = OxmlElement("w:pBdr")
        bottom = OxmlElement("w:bottom")
        bottom.set(qn("w:val"), "single")
        bottom.set(qn("w:sz"), "32")
        bottom.set(qn("w:space"), "1")
        bottom.set(qn("w:color"), "FF0000")
        border.append(bottom)
        p_pr.insert(0, border)
    run = paragraph.add_run(text)
    _set_run(run, font_cn, size_pt, color=color, bold=bold, run_spacing=run_spacing)
    return paragraph


def _set_run(run, font_cn: str, size_pt: int, *, color=BLACK, bold: bool = False, run_spacing: int | None = None) -> None:
    """给一个 run 同时指定西文 Times New Roman 和中文字体。"""
    run.font.name = "Times New Roman"
    run.font.size = Pt(size_pt)
    run.font.bold = bold
    run.font.color.rgb = color
    r_pr = run._element.get_or_add_rPr()
    r_fonts = r_pr.find(qn("w:rFonts"))
    if r_fonts is None:
        r_fonts = OxmlElement("w:rFonts")
        r_pr.insert(0, r_fonts)
    r_fonts.set(qn("w:ascii"), "Times New Roman")
    r_fonts.set(qn("w:hAnsi"), "Times New Roman")
    r_fonts.set(qn("w:eastAsia"), font_cn)
    if run_spacing is not None:
        spacing = OxmlElement("w:spacing")
        spacing.set(qn("w:val"), str(run_spacing))
        r_pr.append(spacing)


def _as_text(value: Any) -> str:
    """把模型可能传入的 list/多行文本收成去空白的字符串。"""
    if value is None:
        return ""
    if isinstance(value, (list, tuple)):
        return "\n\n".join(_as_text(item) for item in value if _as_text(item)).strip()
    return str(value).strip()


def _issuer_from_red_header(red_header: str) -> str:
    """红头「××公司文件」默认落款为「××公司」。"""
    text = red_header.strip()
    if text.endswith("文件"):
        text = text[: -len("文件")].strip()
    return text


def _normalize_issue_date(raw: str) -> str:
    """把 ISO 日期转成「2026年9月10日」；空值用当天。"""
    text = raw.strip()
    if not text:
        today = date.today()
        return f"{today.year}年{today.month}月{today.day}日"
    matched = _ISO_DATE_RE.match(text)
    if matched:
        year, month, day = (int(matched.group(1)), int(matched.group(2)), int(matched.group(3)))
        return f"{year}年{month}月{day}日"
    return text


def _safe_filename(raw: str) -> str:
    """去掉路径分隔符和非法字符，保证落在 outputs/ 文件名层。"""
    stem = Path(raw.strip()).name
    stem = _UNSAFE_FILENAME_RE.sub("_", stem).strip(" .")
    if stem.lower().endswith(".docx"):
        stem = stem[: -len(".docx")]
    stem = stem or "公文"
    stem = stem[:_MAX_FILENAME_STEM]
    return f"{stem}.docx"


def _is_relative_to(path: Path, parent: Path) -> bool:
    """兼容旧 Python 的 Path.is_relative_to。"""
    try:
        path.relative_to(parent)
        return True
    except ValueError:
        return False
