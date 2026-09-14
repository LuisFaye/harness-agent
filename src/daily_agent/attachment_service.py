"""会话附件的本地存储、解析和 Agent 上下文组装。"""

from __future__ import annotations

import csv
import hashlib
import io
import json
import mimetypes
import re
import threading
import uuid
from dataclasses import asdict, dataclass, replace
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

from docx import Document
from openpyxl import load_workbook
from PIL import Image
from pypdf import PdfReader


MAX_ATTACHMENT_BYTES = 20 * 1024 * 1024
MAX_EXTRACTED_CHARS = 60_000
MAX_CONTEXT_CHARS = 80_000
SUPPORTED_SUFFIXES = {
    ".txt", ".md", ".csv", ".json", ".pdf", ".docx", ".xlsx",
    ".png", ".jpg", ".jpeg", ".webp", ".gif",
}
IMAGE_SUFFIXES = {".png", ".jpg", ".jpeg", ".webp", ".gif"}


@dataclass(frozen=True)
class AttachmentRecord:
    id: str
    user_id: str
    session_id: str
    filename: str
    content_type: str
    size: int
    suffix: str
    storage_path: str
    extracted_text: str
    parse_status: str
    parse_message: str
    created_at: str

    def to_public_dict(self, *, include_text: bool = False) -> dict[str, Any]:
        payload = {
            "id": self.id,
            "filename": self.filename,
            "content_type": self.content_type,
            "size": self.size,
            "suffix": self.suffix,
            "parse_status": self.parse_status,
            "parse_message": self.parse_message,
            "created_at": self.created_at,
            "text_preview": self.extracted_text[:500],
            "download_url": f"/attachments/{self.id}/download?session_id={self.session_id}",
        }
        if include_text:
            payload["extracted_text"] = self.extracted_text
        return payload


class AttachmentService:
    """附件保存在用户目录，不接受客户端传入任何真实文件路径。"""

    def __init__(self, root: Path | None = None) -> None:
        self.root = (root or Path.home() / ".sb-agent" / "attachments").resolve()
        self.files_dir = self.root / "files"
        self.records_dir = self.root / "records"
        self.files_dir.mkdir(parents=True, exist_ok=True)
        self.records_dir.mkdir(parents=True, exist_ok=True)
        self._lock = threading.Lock()

    def save(self, *, filename: str, content_type: str, content: bytes, user_id: str, session_id: str) -> AttachmentRecord:
        safe_name = _safe_filename(filename)
        suffix = Path(safe_name).suffix.lower()
        if suffix not in SUPPORTED_SUFFIXES:
            raise ValueError(f"暂不支持 {suffix or '无扩展名'} 文件。支持：{', '.join(sorted(SUPPORTED_SUFFIXES))}")
        if not content:
            raise ValueError("附件内容为空。")
        if len(content) > MAX_ATTACHMENT_BYTES:
            raise ValueError(f"单个附件不能超过 {MAX_ATTACHMENT_BYTES // 1024 // 1024}MB。")

        attachment_id = uuid.uuid4().hex
        namespace = _namespace(user_id, session_id)
        target_dir = self.files_dir / namespace
        target_dir.mkdir(parents=True, exist_ok=True)
        storage_path = target_dir / f"{attachment_id}{suffix}"
        storage_path.write_bytes(content)
        try:
            extracted_text, parse_status, parse_message = _extract_content(storage_path, suffix)
        except Exception as exc:
            extracted_text, parse_status, parse_message = "", "error", f"解析失败：{exc}"

        record = AttachmentRecord(
            id=attachment_id,
            user_id=str(user_id or "default"),
            session_id=str(session_id or "work"),
            filename=safe_name,
            content_type=content_type or mimetypes.guess_type(safe_name)[0] or "application/octet-stream",
            size=len(content),
            suffix=suffix,
            storage_path=str(storage_path),
            extracted_text=extracted_text[:MAX_EXTRACTED_CHARS],
            parse_status=parse_status,
            parse_message=parse_message,
            created_at=datetime.now(timezone.utc).isoformat(),
        )
        self._write_record(record)
        return record

    def get(self, attachment_id: str, *, user_id: str, session_id: str | None = None) -> AttachmentRecord:
        record = self._read_record(attachment_id)
        if record.user_id != user_id or session_id is not None and record.session_id != session_id:
            raise ValueError("附件不存在或不属于当前会话。")
        return record

    def list_session(self, *, user_id: str, session_id: str) -> list[AttachmentRecord]:
        records = []
        for path in self.records_dir.glob("*.json"):
            try:
                record = AttachmentRecord(**json.loads(path.read_text(encoding="utf-8")))
            except (OSError, TypeError, ValueError, json.JSONDecodeError):
                continue
            if record.user_id == user_id and record.session_id == session_id:
                records.append(record)
        return sorted(records, key=lambda item: item.created_at)

    def build_context(
        self,
        attachment_ids: list[str],
        *,
        user_id: str,
        session_id: str,
        vision_enabled: bool = False,
    ) -> str:
        if not attachment_ids:
            return ""
        sections = [
            "本轮附件内容（只能依据这里真实提取到的内容回答，不得假装读取解析失败或不支持识别的部分）："
        ]
        if vision_enabled:
            sections.append("图片附件已直接传入视觉模型，请结合用户问题理解图片内容。")
        used_chars = 0
        for attachment_id in list(dict.fromkeys(attachment_ids))[:10]:
            record = self.get(attachment_id, user_id=user_id, session_id=session_id)
            if vision_enabled and record.suffix in IMAGE_SUFFIXES:
                continue
            remaining = MAX_CONTEXT_CHARS - used_chars
            if remaining <= 0:
                sections.append("- 其余附件因上下文长度限制未注入。")
                break
            content = record.extracted_text[:remaining] if record.extracted_text else f"[{record.parse_message}]"
            sections.append(f"\n## 附件：{record.filename}\n解析状态：{record.parse_status}\n{content}")
            used_chars += len(content)
        return "\n".join(sections)

    def build_vision_images(
        self,
        attachment_ids: list[str],
        *,
        user_id: str,
        session_id: str,
    ) -> list[Any]:
        """把图片附件转成 Agno Image，供视觉模型直接理解。"""
        try:
            from agno.media import Image as AgnoImage
        except ImportError:
            return []

        images: list[Any] = []
        for attachment_id in list(dict.fromkeys(attachment_ids))[:10]:
            record = self.get(attachment_id, user_id=user_id, session_id=session_id)
            if record.suffix not in IMAGE_SUFFIXES:
                continue
            path = Path(record.storage_path)
            if not path.is_file():
                continue
            images.append(
                AgnoImage(
                    content=path.read_bytes(),
                    mime_type=record.content_type or mimetypes.guess_type(record.filename)[0],
                    format=record.suffix.lstrip(".") or None,
                    detail="auto",
                )
            )
        return images

    def claim_legacy_user(self, old_user_id: str, new_user_id: str) -> int:
        """首位所有者接管历史附件，同时迁移按用户派生的内容目录。"""
        changed = 0
        with self._lock:
            for record_path in self.records_dir.glob("*.json"):
                try:
                    record = AttachmentRecord(**json.loads(record_path.read_text(encoding="utf-8")))
                    if record.user_id != old_user_id:
                        continue
                    expected_source = (
                        self.files_dir / _namespace(old_user_id, record.session_id) / f"{record.id}{record.suffix}"
                    ).resolve()
                    if Path(record.storage_path).resolve() != expected_source or not expected_source.is_file():
                        continue
                    target = (
                        self.files_dir / _namespace(new_user_id, record.session_id) / f"{record.id}{record.suffix}"
                    ).resolve()
                    target.parent.mkdir(parents=True, exist_ok=True)
                    expected_source.replace(target)
                    migrated = replace(record, user_id=new_user_id, storage_path=str(target))
                    record_path.write_text(
                        json.dumps(asdict(migrated), ensure_ascii=False, indent=2),
                        encoding="utf-8",
                    )
                    changed += 1
                except (OSError, TypeError, ValueError, json.JSONDecodeError):
                    continue
        return changed

    def _write_record(self, record: AttachmentRecord) -> None:
        path = self.records_dir / f"{record.id}.json"
        with self._lock:
            path.write_text(json.dumps(asdict(record), ensure_ascii=False, indent=2), encoding="utf-8")

    def _read_record(self, attachment_id: str) -> AttachmentRecord:
        if not re.fullmatch(r"[a-f0-9]{32}", str(attachment_id or "")):
            raise ValueError("附件 ID 不合法。")
        path = self.records_dir / f"{attachment_id}.json"
        if not path.exists():
            raise ValueError("附件不存在。")
        return AttachmentRecord(**json.loads(path.read_text(encoding="utf-8")))


def _safe_filename(filename: str) -> str:
    value = Path(str(filename or "attachment")).name.strip().replace("\x00", "")
    return value[:180] or "attachment"


def _namespace(user_id: str, session_id: str) -> str:
    raw = f"{user_id}\x00{session_id}".encode("utf-8", errors="replace")
    return hashlib.sha256(raw).hexdigest()[:24]


def _extract_content(path: Path, suffix: str) -> tuple[str, str, str]:
    if suffix in {".txt", ".md", ".json"}:
        text = _decode_text(path.read_bytes())
        return text, "success", f"已提取 {len(text)} 个字符。"
    if suffix == ".csv":
        return _extract_csv(path)
    if suffix == ".pdf":
        return _extract_pdf(path)
    if suffix == ".docx":
        return _extract_docx(path)
    if suffix == ".xlsx":
        return _extract_xlsx(path)
    if suffix in IMAGE_SUFFIXES:
        return _extract_image_metadata(path)
    raise ValueError("不支持的附件类型。")


def _decode_text(content: bytes) -> str:
    for encoding in ("utf-8-sig", "utf-8", "gb18030"):
        try:
            return content.decode(encoding)
        except UnicodeDecodeError:
            continue
    return content.decode("utf-8", errors="replace")


def _extract_csv(path: Path) -> tuple[str, str, str]:
    text = _decode_text(path.read_bytes())
    rows = list(csv.reader(io.StringIO(text)))[:1000]
    rendered = "\n".join(" | ".join(cell.strip() for cell in row) for row in rows)
    return rendered, "success", f"已读取 {len(rows)} 行 CSV。"


def _extract_pdf(path: Path) -> tuple[str, str, str]:
    reader = PdfReader(str(path))
    pages = []
    for index, page in enumerate(reader.pages[:200], start=1):
        pages.append(f"[第 {index} 页]\n{page.extract_text() or ''}")
    text = "\n\n".join(pages).strip()
    status = "success" if text else "partial"
    message = f"已读取 {len(pages)} 页 PDF。" if text else "PDF 没有可提取文字，可能是扫描件。"
    return text, status, message


def _extract_docx(path: Path) -> tuple[str, str, str]:
    document = Document(str(path))
    lines = [paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()]
    for table in document.tables:
        for row in table.rows:
            lines.append(" | ".join(cell.text.strip() for cell in row.cells))
    text = "\n".join(lines)
    return text, "success" if text else "partial", f"已提取 {len(lines)} 段 Word 内容。"


def _extract_xlsx(path: Path) -> tuple[str, str, str]:
    workbook = load_workbook(str(path), read_only=True, data_only=True)
    lines = []
    row_count = 0
    for sheet in workbook.worksheets[:20]:
        lines.append(f"[工作表：{sheet.title}]")
        for row in sheet.iter_rows(values_only=True):
            values = ["" if value is None else str(value) for value in row]
            if any(values):
                lines.append(" | ".join(values))
                row_count += 1
            if row_count >= 2000:
                break
        if row_count >= 2000:
            break
    workbook.close()
    return "\n".join(lines), "success", f"已读取 {len(workbook.sheetnames)} 个工作表中的 {row_count} 行。"


def _extract_image_metadata(path: Path) -> tuple[str, str, str]:
    with Image.open(path) as image:
        text = (
            f"图片格式：{image.format or path.suffix.lstrip('.')}\n"
            f"图片尺寸：{image.width} × {image.height}\n"
            f"颜色模式：{image.mode}"
        )
    return text, "success", "图片已上传；使用视觉模型时会直接识别内容。"
