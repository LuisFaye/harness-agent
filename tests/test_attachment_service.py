from pathlib import Path

import pytest
from docx import Document
from openpyxl import Workbook
from PIL import Image

from daily_agent.attachment_service import AttachmentService


def test_text_attachment_is_session_scoped_and_builds_agent_context(tmp_path):
    service = AttachmentService(root=tmp_path / "attachments")

    record = service.save(
        filename="需求.md",
        content_type="text/markdown",
        content="项目代号是星河，预算为 42 万元。".encode(),
        user_id="u1",
        session_id="s1",
    )

    assert record.parse_status == "success"
    assert "星河" in record.extracted_text
    assert service.list_session(user_id="u1", session_id="s1") == [record]
    assert "预算为 42 万元" in service.build_context([record.id], user_id="u1", session_id="s1")
    with pytest.raises(ValueError, match="不属于当前会话"):
        service.get(record.id, user_id="u1", session_id="s2")


def test_docx_and_xlsx_are_parsed(tmp_path):
    service = AttachmentService(root=tmp_path / "attachments")
    docx_path = tmp_path / "report.docx"
    document = Document()
    document.add_paragraph("Word 中的季度结论")
    document.save(docx_path)
    xlsx_path = tmp_path / "data.xlsx"
    workbook = Workbook()
    workbook.active.append(["指标", "数值"])
    workbook.active.append(["收入", 123])
    workbook.save(xlsx_path)

    docx_record = service.save(
        filename=docx_path.name,
        content_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        content=docx_path.read_bytes(),
        user_id="u1",
        session_id="s1",
    )
    xlsx_record = service.save(
        filename=xlsx_path.name,
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        content=xlsx_path.read_bytes(),
        user_id="u1",
        session_id="s1",
    )

    assert "季度结论" in docx_record.extracted_text
    assert "收入 | 123" in xlsx_record.extracted_text


def test_image_reports_metadata_without_faking_ocr(tmp_path):
    service = AttachmentService(root=tmp_path / "attachments")
    image_path = tmp_path / "sample.png"
    Image.new("RGB", (64, 32), "white").save(image_path)

    record = service.save(
        filename=image_path.name,
        content_type="image/png",
        content=image_path.read_bytes(),
        user_id="u1",
        session_id="s1",
    )

    assert record.parse_status == "success"
    assert "64 × 32" in record.extracted_text
    assert "视觉模型" in record.parse_message


def test_build_vision_images_returns_agno_image_payload(tmp_path):
    service = AttachmentService(root=tmp_path / "attachments")
    image_path = tmp_path / "sample.png"
    Image.new("RGB", (32, 16), "red").save(image_path)

    record = service.save(
        filename=image_path.name,
        content_type="image/png",
        content=image_path.read_bytes(),
        user_id="u1",
        session_id="s1",
    )
    text_record = service.save(
        filename="note.txt",
        content_type="text/plain",
        content=b"hello",
        user_id="u1",
        session_id="s1",
    )

    images = service.build_vision_images([record.id, text_record.id], user_id="u1", session_id="s1")

    assert len(images) == 1
    assert images[0].content
    assert images[0].mime_type == "image/png"


def test_build_context_skips_images_when_vision_enabled(tmp_path):
    service = AttachmentService(root=tmp_path / "attachments")
    image_path = tmp_path / "sample.png"
    Image.new("RGB", (8, 8), "blue").save(image_path)
    record = service.save(
        filename=image_path.name,
        content_type="image/png",
        content=image_path.read_bytes(),
        user_id="u1",
        session_id="s1",
    )

    context = service.build_context([record.id], user_id="u1", session_id="s1", vision_enabled=True)

    assert "视觉模型" in context
    assert "图片尺寸" not in context


def test_unsafe_filename_and_unsupported_extension_are_rejected(tmp_path):
    service = AttachmentService(root=tmp_path / "attachments")

    record = service.save(
        filename="../../safe.txt",
        content_type="text/plain",
        content=b"safe",
        user_id="u1",
        session_id="s1",
    )

    assert record.filename == "safe.txt"
    assert Path(record.storage_path).is_relative_to(service.files_dir)
    with pytest.raises(ValueError, match="暂不支持"):
        service.save(
            filename="bad.exe",
            content_type="application/octet-stream",
            content=b"bad",
            user_id="u1",
            session_id="s1",
        )
